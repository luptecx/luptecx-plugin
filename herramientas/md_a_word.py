# /// script
# requires-python = ">=3.12"
# dependencies = ["python-docx==1.2.0", "docx2pdf==0.1.8"]
# ///
"""Arma un documento de LUPTECX en Word (y, si se pide, en PDF) desde uno o más archivos
Markdown: el contrato, el acta de entrega, el manual del cliente.

    uv run --script md_a_word.py --salida documento.docx [--datos datos.json]
        [--encabezado TEXTO] [--pdf] parte1.md [parte2.md …]

Cada parte empieza en una página nueva. Cada ``{{campo}}`` se reemplaza con su valor de
``datos.json``; si falta uno, el script se niega y dice cuáles (un dato que no se sabe
va como ``[POR CONFIRMAR: …]``, nunca vacío). Lo pendiente ([POR DEFINIR …] y
[POR CONFIRMAR …]) sale resaltado en amarillo. El estilo sigue `luptecx-estilo-visual`
(negro sobre blanco, dorado profundo solo en el encabezado, tablas con fondo crema).

Markdown que entiende: ``#``, ``##`` y ``###``, viñetas ``-``, listas ``1.``, tablas,
``**negrilla**`` y citas ``>``. Las líneas seguidas forman un solo párrafo.
"""

import argparse
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

CAMPO = re.compile(r"\{\{(\w+)\}\}")
NEGRILLA_O_PENDIENTE = re.compile(
    r"(\*\*[^*]+\*\*|`[^`]+`|\[POR (?:DEFINIR|CONFIRMAR)[^\]]*\])"
)
PENDIENTE = re.compile(r"\[POR (?:DEFINIR|CONFIRMAR)")
FUENTE = "Segoe UI"  # Poppins si está instalada; Segoe UI es la alternativa del estilo
NEGRO = RGBColor(0x14, 0x14, 0x14)
GRIS = RGBColor(0x5E, 0x5A, 0x52)
DORADO_PROFUNDO = RGBColor(0x86, 0x66, 0x0C)
CREMA = "F7F3E8"


def llenar(texto: str, datos: dict[str, str]) -> tuple[str, set[str]]:
    faltan = {c for c in CAMPO.findall(texto) if not str(datos.get(c, "")).strip()}
    return CAMPO.sub(lambda m: str(datos.get(m.group(1), m.group(0))), texto), faltan


def agregar_texto(parrafo, texto: str) -> None:
    for trozo in NEGRILLA_O_PENDIENTE.split(texto):
        if not trozo:
            continue
        if trozo.startswith("**"):
            parrafo.add_run(trozo[2:-2]).bold = True
        elif trozo.startswith("`"):
            parrafo.add_run(trozo[1:-1]).font.name = "Consolas"
        elif trozo.startswith("[POR "):
            parrafo.add_run(trozo).font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            parrafo.add_run(trozo)


MARCAS = (
    ("# ", "titulo"),
    ("## ", "h1"),
    ("### ", "h2"),
    ("- ", "vineta"),
)


def bloques(md: str) -> list[tuple[str, str | list[list[str]]]]:
    """Título, subtítulos, viñetas, listas numeradas, citas, tablas y párrafos. Las
    líneas seguidas (y las sangradas bajo una viñeta) se unen en un solo bloque."""
    salida: list[tuple[str, str | list[list[str]]]] = []
    for linea in md.splitlines():
        if not linea.strip():
            salida.append(("vacio", ""))
            continue
        if linea.startswith("|"):
            celdas = [c.strip() for c in linea.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-+:?", c) for c in celdas):
                continue
            ultimo = salida[-1] if salida else None
            if ultimo and ultimo[0] == "tabla" and isinstance(ultimo[1], list):
                ultimo[1].append(celdas)
            else:
                salida.append(("tabla", [celdas]))
            continue
        if linea.startswith(">"):
            cita = linea.lstrip("> ").strip()
            if salida and salida[-1][0] == "cita":
                salida[-1] = ("cita", f"{salida[-1][1]} {cita}".strip())
            elif cita:
                salida.append(("cita", cita))
            continue
        for marca, tipo in MARCAS:
            if linea.startswith(marca):
                salida.append((tipo, linea[len(marca) :].strip()))
                break
        else:
            if re.match(r"\d+\. ", linea):
                salida.append(("numerada", linea.strip()))
                continue
            anterior = salida[-1][0] if salida else ""
            sangrada = linea.startswith(" ") and anterior in ("vineta", "numerada")
            if sangrada or anterior == "parrafo":
                tipo, texto = salida[-1]
                salida[-1] = (tipo, f"{texto} {linea.strip()}")
            else:
                salida.append(("parrafo", linea.strip()))
    return [b for b in salida if b[0] != "vacio"]


def sombrear(celda, color: str) -> None:
    propiedades = celda._tc.get_or_add_tcPr()
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:fill"), color)
    propiedades.append(sombra)


def escribir(documento, md: str) -> None:
    for tipo, contenido in bloques(md):
        if tipo == "tabla" and isinstance(contenido, list):
            tabla = documento.add_table(rows=len(contenido), cols=len(contenido[0]))
            tabla.style = "Table Grid"
            for i, fila in enumerate(contenido):
                for j, celda in enumerate(fila):
                    parrafo = tabla.cell(i, j).paragraphs[0]
                    agregar_texto(parrafo, celda)
                    if i == 0:
                        sombrear(tabla.cell(i, j), CREMA)
                        for run in parrafo.runs:
                            run.bold = True
            documento.add_paragraph()
            continue
        texto = str(contenido)
        if tipo == "titulo":
            documento.add_heading(texto, level=0)
        elif tipo in ("h1", "h2"):
            documento.add_heading(texto, level=1 if tipo == "h1" else 2)
        elif tipo == "vineta":
            agregar_texto(documento.add_paragraph(style="List Bullet"), texto)
        elif tipo == "cita":
            parrafo = documento.add_paragraph()
            agregar_texto(parrafo, texto)
            for run in parrafo.runs:
                run.font.color.rgb = GRIS
        else:
            agregar_texto(documento.add_paragraph(), texto)


def dar_estilo(documento) -> None:
    for nombre, tamano in (
        ("Normal", 11),
        ("Title", 24),
        ("Heading 1", 15),
        ("Heading 2", 12.5),
    ):
        estilo = documento.styles[nombre]
        estilo.font.name = FUENTE
        estilo.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
        estilo.font.size = Pt(tamano)
        estilo.font.color.rgb = NEGRO
        if nombre != "Normal":
            estilo.font.bold = True


def main() -> int:
    argumentos = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    argumentos.add_argument("partes", nargs="+", type=Path)
    argumentos.add_argument("--salida", type=Path, required=True)
    argumentos.add_argument("--datos", type=Path)
    argumentos.add_argument("--encabezado", default="")
    argumentos.add_argument("--pdf", action="store_true")
    opciones = argumentos.parse_args()

    datos = {}
    if opciones.datos:
        datos = json.loads(opciones.datos.read_text(encoding="utf-8"))
    textos, faltan = [], set()
    for parte in opciones.partes:
        texto, sin_valor = llenar(parte.read_text(encoding="utf-8"), datos)
        textos.append(texto)
        faltan |= sin_valor
    if faltan:
        print("Faltan datos (usa [POR CONFIRMAR: …] si no se saben):")
        print(", ".join(sorted(faltan)))
        return 1

    documento = Document()
    dar_estilo(documento)
    if opciones.encabezado:
        run = documento.sections[0].header.paragraphs[0].add_run(opciones.encabezado)
        run.bold = True
        run.font.color.rgb = DORADO_PROFUNDO
    for i, texto in enumerate(textos):
        if i:
            documento.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        escribir(documento, texto)
    documento.save(opciones.salida)
    pendientes = sum(len(PENDIENTE.findall(t)) for t in textos)
    print(f"Word: {opciones.salida} · {pendientes} pendientes resaltados")

    if opciones.pdf:
        pdf = opciones.salida.with_suffix(".pdf")
        try:
            from docx2pdf import convert  # noqa: PLC0415 — solo hace falta con --pdf

            convert(str(opciones.salida), str(pdf))
        except Exception as error:  # noqa: BLE001 — sin Word instalado, se dice y sigue
            print(
                f"No se pudo pasar a PDF ({error}). Ábralo en Word y guárdelo en PDF."
            )
            return 0
        print(f"PDF: {pdf}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
